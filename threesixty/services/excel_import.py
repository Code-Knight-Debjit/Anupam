"""Customer import from Excel (REQ-IMP-01).

Rows sharing a company name (case-insensitive) are one customer with one
contact per row; existing customers are updated, blank cells keep what is
saved; header names tolerate case and spacing; map pins accept "lat, lng" or
a Google Maps link. Returns a summary and a list of row-level warnings.
"""
import io

from django.contrib.auth import get_user_model
from django.db import transaction

from ..models import Contact, Customer, LookupItem
from ..rules.anupam_rules import title_case_words
from .errors import ServiceError

User = get_user_model()

TEMPLATE_COLS = [
    'Company', 'Plant', 'Address', 'Area', 'City', 'State', 'Pincode', 'Website', 'GST Number', 'Plant Location',
    'Head Office Name', 'Head Office Address', 'Head Office City', 'Head Office State', 'Head Office Pincode',
    'Head Office Phone', 'Head Office Location', 'Industry', 'Source', 'Potential', 'Transport Preference',
    'Courier Preference', 'Preferred Languages', 'Assigned To', 'Interested Brands', 'Interested Products',
    'Contact Name', 'Job Title', 'Phone', 'Phone 2', 'Email',
]
SAMPLE_ROW = {
    'Company': 'Sundaram Fasteners Ltd', 'Plant': 'Unit 2', 'Address': 'Plot 12, Industrial Estate', 'Area': 'Peenya',
    'City': 'Bengaluru', 'State': 'Karnataka', 'Pincode': '560058', 'Website': 'www.sundaramfasteners.com',
    'GST Number': '29ABCDE1234F1Z5', 'Plant Location': '13.0285, 77.5197', 'Head Office Name': 'Corporate Office',
    'Head Office City': 'Chennai', 'Head Office State': 'Tamil Nadu', 'Industry': 'Gearbox Manufacturing',
    'Source': 'Referral', 'Potential': 'High', 'Preferred Languages': 'English, Kannada',
    'Interested Brands': 'Timken, SKF', 'Interested Products': 'Spherical Roller Bearing',
    'Contact Name': 'Ramesh Iyer', 'Job Title': 'Purchase', 'Phone': '9876543210', 'Email': 'ramesh@example.com',
}
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
_CANON = {' '.join(c.lower().split()): c for c in TEMPLATE_COLS}


def template_workbook_bytes():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Template'
    ws.append(TEMPLATE_COLS)
    ws.append([SAMPLE_ROW.get(c, '') for c in TEMPLATE_COLS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_from_workbook(uploaded):
    """Rows of dicts keyed by the canonical column names (unknown columns ignored)."""
    from openpyxl import load_workbook
    if uploaded.size > MAX_BYTES:
        raise ServiceError(400, 'That file is too large (max 5 MB).')
    try:
        wb = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception:
        raise ServiceError(400, 'That file is not a valid .xlsx workbook.')
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise ServiceError(400, 'The sheet is empty.')
    columns = [_CANON.get(' '.join(str(h or '').lower().split())) for h in header]
    if 'Company' not in columns:
        raise ServiceError(400, 'The sheet needs a "Company" column.')
    rows = []
    for values in rows_iter:
        rows.append({col: ('' if v is None else str(v).strip()) for col, v in zip(columns, values) if col})
        if len(rows) > MAX_ROWS:
            raise ServiceError(400, f'Too many rows (max {MAX_ROWS}).')
    return rows


def _split(cell):
    return [x.strip() for x in str(cell).split(',') if x.strip()]


def _lookup_ids(kind, names):
    ids = []
    for name in names:
        item, _ = LookupItem.objects.get_or_create(kind=kind, name=name)
        if not item.is_active:
            item.is_active = True
            item.save(update_fields=['is_active'])
        ids.append(item.pk)
    return ids


def _user_for(name):
    name = name.strip().lower()
    for u in User.objects.filter(is_active=True, threesixty_profile__isnull=False):
        if name in (u.username.lower(), (u.get_full_name() or '').lower()):
            return u
    return None


@transaction.atomic
def import_rows(rows, user):
    from . import customers as customer_service
    new = updated = contacts_added = 0
    warnings = []
    by_name = {c.name.strip().lower(): c for c in Customer.objects.all()}
    for number, row in enumerate(rows, start=2):
        name = title_case_words(str(row.get('Company') or '').strip())
        if not name:
            continue
        customer = by_name.get(name.lower())
        if customer is None:
            customer = Customer(name=name, created_by=user,
                                created_by_role=user.threesixty_profile.role if hasattr(user, 'threesixty_profile') else '')
            created = True
        else:
            created = False
        data = {}
        for col, field in (('Plant', 'plant'), ('Area', 'area'), ('City', 'city'), ('State', 'state'), ('Pincode', 'pincode'),
                           ('Website', 'website'), ('Industry', 'industry'), ('Source', 'source'), ('Address', 'address'),
                           ('GST Number', 'gst_number'), ('Head Office Name', 'ho_name'),
                           ('Head Office Address', 'ho_address'), ('Head Office City', 'ho_city'),
                           ('Head Office State', 'ho_state'), ('Head Office Pincode', 'ho_pincode'),
                           ('Head Office Phone', 'ho_phone'), ('Transport Preference', 'transport_preference'),
                           ('Courier Preference', 'courier_preference')):
            if row.get(col):
                data[field] = str(row[col])
        potential = {'high': 'High', 'medium': 'Medium', 'low': 'Low'}.get(str(row.get('Potential') or '').strip().lower())
        if potential:
            data['potential'] = potential
        elif row.get('Potential'):
            warnings.append(f'Row {number}: potential "{row["Potential"]}" ignored (use High, Medium or Low).')
        for col, key in (('Plant Location', 'plant_geo'), ('Head Office Location', 'ho_geo')):
            if row.get(col):
                try:
                    customer_service.parse_pin(row[col])
                    data[key] = row[col]
                except ServiceError:
                    warnings.append(f'Row {number}: {col} "{row[col]}" is not a valid location — ignored.')
        customer_service.apply_customer_data(customer, data, user)
        customer.save()
        if created:
            new += 1
            by_name[name.lower()] = customer
        else:
            updated += 1                      # counted per row, as the reference app does
        m2m = {}
        for col, key, kind in (('Preferred Languages', 'preferred_language_ids', LookupItem.KIND_LANGUAGE),
                               ('Interested Brands', 'interested_brand_ids', LookupItem.KIND_BRAND),
                               ('Interested Products', 'interested_product_ids', LookupItem.KIND_PRODUCT)):
            if row.get(col):
                m2m[key] = _lookup_ids(kind, _split(row[col]))
        customer_service.apply_customer_m2m(customer, m2m)
        if row.get('Assigned To'):
            for person in _split(row['Assigned To']):
                assignee = _user_for(person)
                if assignee is None:
                    warnings.append(f'Row {number}: user "{person}" not found — not assigned.')
                else:
                    customer.assigned_to.add(assignee)
        cname = title_case_words(str(row.get('Contact Name') or '').strip())
        phone = str(row.get('Phone') or '').strip()
        if cname or phone or row.get('Phone 2') or row.get('Email'):
            if not customer.contacts.filter(name__iexact=cname, phone=phone).exists():
                Contact.objects.create(
                    customer=customer, name=cname, job_title=str(row.get('Job Title') or '').strip(), phone=phone,
                    phone2=str(row.get('Phone 2') or '').strip(), email=str(row.get('Email') or '').strip())
                contacts_added += 1
    return {'new': new, 'updated': updated, 'contacts': contacts_added, 'warnings': warnings}
