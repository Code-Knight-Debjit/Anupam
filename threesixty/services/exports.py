"""Excel and PDF export of tables (lists, history pages, reports) and of a
single sales document. Cell text starting with = + - @ is neutralised so a
customer name like "=HYPERLINK(...)" can't run as a spreadsheet formula."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from .errors import ServiceError
from .printing import money, print_totals

_FORMULA_START = ('=', '+', '-', '@', '\t', '\r')


def safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def table_xlsx(title, headers, rows) -> bytes:
    if not rows:
        raise ServiceError(400, 'Nothing to export.')
    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Sheet')[:31]
    ws.append([safe_cell(h) for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    for column in ws.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(longest + 2, 50)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def table_pdf(title, headers, rows, compress=True) -> bytes:
    if not rows:
        raise ServiceError(400, 'Nothing to export.')
    small = ParagraphStyle('s', fontName='Helvetica', fontSize=7, leading=8.5)
    head = ParagraphStyle('h', parent=small, fontName='Helvetica-Bold', textColor=colors.white)
    cell = lambda v, style=small: Paragraph(escape('' if v is None else str(v)), style)
    data = [[cell(h, head) for h in headers]] + [[cell(v) for v in row] for row in rows]
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=title, pageCompression=1 if compress else 0)
    table = Table(data, repeatRows=1, style=TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#a14000')), ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f5f7')])]))
    big = ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=12)
    doc.build([Paragraph(escape(title), big), Spacer(1, 3 * mm), table])
    return out.getvalue()


def document_xlsx(document) -> bytes:
    """One document, as printed (printed totals, not the entry-form totals)."""
    totals = print_totals(document)
    headers = ['Sr', 'Description', 'Qty', 'Unit', 'Rate', 'Amount']
    rows = [[l['sr'], l['name'], float(l['qty']), l['unit'], float(l['rate']), float(money(l['amount']))] for l in totals['lines']]
    wb = Workbook()
    ws = wb.active
    ws.title = document.get_type_display()[:31]
    ws.append([safe_cell(f'{document.get_type_display()} {document.doc_number}')])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Customer', safe_cell(document.customer.select_label)])
    ws.append(['Date', document.date.isoformat()])
    ws.append(['Branch', safe_cell(document.branch.name)])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    ws.append([])
    for label, key in (('Sub Total', 'sub_total'), ('Tax Amount', 'tax_total'), ('Packing', 'packing'), ('Freight', 'freight'),
                       ('Other charges', 'other'), ('Round off', 'round_off'), ('TOTAL', 'total')):
        ws.append(['', label, '', '', '', float(money(totals[key]))])
    ws.column_dimensions['B'].width = 46
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
